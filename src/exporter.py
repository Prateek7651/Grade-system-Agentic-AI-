import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def _autofit_and_style(filepath, sheet_name='Sheet1'):
    import openpyxl
    wb = openpyxl.load_workbook(filepath)
    ws = wb[sheet_name]
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    for cell in ws[1]:
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        max_len = max(len(str(c.value)) if c.value else 0 for c in column_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)
    ws.freeze_panes = 'A2'
    wb.save(filepath)

def export_dataframe(df, name, output_dir='data/processed/'):
    out_path = Path(output_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    csv_path = out_path / f'{name}.csv'
    xlsx_path = out_path / f'{name}.xlsx'
    df.to_csv(csv_path, index=False)
    df.to_excel(xlsx_path, index=False, sheet_name='Sheet1')
    _autofit_and_style(str(xlsx_path))
    print(f"  Saved: {csv_path}")
    print(f"  Saved: {xlsx_path}")
    return {'csv': str(csv_path), 'xlsx': str(xlsx_path)}

def export_all_outputs(results, output_dir='data/processed/'):
    paths = {}
    print("Exporting Master Performance File...")
    master_cols = ['rank','name','email','quizzes_attempted','total_marks_scored','total_marks_possible','avg_percentage','final_percentile','grade']
    paths['master'] = export_dataframe(results['master'][master_cols], 'master_performance', output_dir)
    print("\nExporting Module Summary File...")
    module_cols = ['name','email','module','marks_scored','marks_possible','module_percentage','module_percentile']
    module_sorted = results['module'].sort_values(['module','module_percentage'], ascending=[True,False])
    paths['module'] = export_dataframe(module_sorted[module_cols], 'module_summary', output_dir)
    print("\nExporting Final Rankings File...")
    paths['rankings'] = export_dataframe(results['master'][['rank','name','email','avg_percentage','grade']], 'final_rankings', output_dir)
    return paths
